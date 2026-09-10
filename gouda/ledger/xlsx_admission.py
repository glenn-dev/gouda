"""Bounded, non-semantic admission for private browser-uploaded XLSX bytes."""

from __future__ import annotations

from dataclasses import dataclass, field
from io import BytesIO
from pathlib import PurePosixPath
import re
import posixpath
from zipfile import BadZipFile, ZIP_DEFLATED, ZIP_STORED, ZipFile

from defusedxml import ElementTree as SafeElementTree
from defusedxml.common import DefusedXmlException
from openpyxl.styles.numbers import BUILTIN_FORMATS


MAX_ARCHIVE_ENTRIES = 128
MAX_TOTAL_EXPANDED_BYTES = 20 * 1024 * 1024
MAX_ENTRY_EXPANDED_BYTES = 10 * 1024 * 1024
MAX_WORKSHEETS = 4
MAX_WORKSHEET_ROWS = 10_000
MAX_WORKSHEET_COLUMNS = 64
MAX_RECTANGULAR_CELLS = 100_000
MAX_XML_DEPTH = 64
MAX_XML_ELEMENTS = 500_000
MAX_XML_NAME_LENGTH = 1024

# These package relationships do not invoke secondary object/image/pivot
# loaders. Unsupported features are rejected, never stripped from the artifact.
_SAFE_RELATIONSHIPS = frozenset({
    "http://schemas.openxmlformats.org/officeDocument/2006/relationships/" + name
    for name in ("officeDocument", "worksheet", "styles", "sharedStrings", "theme", "extended-properties", "custom-properties")
}) | {"http://schemas.openxmlformats.org/package/2006/relationships/metadata/core-properties"}


@dataclass
class _PackageReferences:
    sheets: list[str] = field(default_factory=list)
    worksheet_targets: list[str] = field(default_factory=list)
    workbook_sheets: dict[str, list[str]] = field(default_factory=dict)
    relationships: dict[str, dict[str, tuple[str, str]]] = field(default_factory=dict)
    shared_strings: list[int] = field(default_factory=list)
    shared_uses: dict[int, int] = field(default_factory=dict)
    style_uses: dict[int, int] = field(default_factory=dict)
    formats: dict[int, int] = field(default_factory=dict)
    styles: list[int] = field(default_factory=list)
    inline_text: int = 0
    string_tables: int = 0
    style_tables: int = 0

    def validate(self):
        if len(self.sheets) > MAX_WORKSHEETS:
            raise XlsxAdmissionError("statement_resource_limit")
        if (len(set(self.sheets)) != len(self.sheets)
                or len(set(self.worksheet_targets)) != len(self.worksheet_targets)):
            raise XlsxAdmissionError("xlsx_invalid")
        used = set()
        for workbook, ids in self.workbook_sheets.items():
            rels = posixpath.join(posixpath.dirname(workbook), "_rels", posixpath.basename(workbook) + ".rels")
            for identifier in ids:
                kind, target = self.relationships.get(rels, {}).get(identifier, ("", ""))
                if not kind.endswith("/worksheet") or target in used:
                    raise XlsxAdmissionError("xlsx_invalid")
                used.add(target)
        if self.string_tables > 1 or self.style_tables > 1:
            raise XlsxAdmissionError("xlsx_invalid")
        # A tiny shared string/style table can expand into gigabytes of copied
        # parser text and serialized row evidence. Account for every cell use.
        text_size = self.inline_text
        for index, count in self.shared_uses.items():
            if index >= len(self.shared_strings):
                raise XlsxAdmissionError("xlsx_invalid")
            text_size += self.shared_strings[index] * count
        for index, count in self.style_uses.items():
            if index >= len(self.styles):
                if index == 0 and not self.styles:
                    continue
                raise XlsxAdmissionError("xlsx_invalid")
            format_id = self.styles[index]
            text_size += self.formats.get(format_id, len(BUILTIN_FORMATS.get(format_id, "General"))) * count
        if text_size * 4 > MAX_TOTAL_EXPANDED_BYTES:
            raise XlsxAdmissionError("statement_resource_limit")

_CELL_REFERENCE = re.compile(r"^\$?([A-Z]{1,3})\$?([1-9][0-9]*)$")
_RANGE_REFERENCE = re.compile(
    r"^\$?([A-Z]{1,3})\$?([1-9][0-9]*)(?::\$?([A-Z]{1,3})\$?([1-9][0-9]*))?$"
)


class XlsxAdmissionError(ValueError):
    def __init__(self, code: str):
        self.code = code
        super().__init__(code)


@dataclass
class _SheetExtent:
    max_row: int = 0
    max_column: int = 0
    inferred_row: int = 0
    inferred_column: int = 0

    def include(self, row: int, column: int) -> None:
        if row > MAX_WORKSHEET_ROWS or column > MAX_WORKSHEET_COLUMNS:
            raise XlsxAdmissionError("statement_resource_limit")
        self.max_row = max(self.max_row, row)
        self.max_column = max(self.max_column, column)


class _BoundedEntryReader:
    """Count actual decompression while exposing only the file interface used by XML."""

    def __init__(self, source, *, total_counter: list[int]):
        self._source = source
        self._entry_count = 0
        self._total_counter = total_counter

    def read(self, size=-1):
        chunk = self._source.read(size)
        self._entry_count += len(chunk)
        self._total_counter[0] += len(chunk)
        if (
            self._entry_count > MAX_ENTRY_EXPANDED_BYTES
            or self._total_counter[0] > MAX_TOTAL_EXPANDED_BYTES
        ):
            raise XlsxAdmissionError("statement_resource_limit")
        return chunk

    def close(self):
        self._source.close()


class _XmlBudgetTarget:
    """Check parser callbacks before iterparse can queue expanded element names."""

    def __init__(self, counter):
        self.counter = counter
        self.depth = 0

    def start_ns(self, prefix, uri):
        if len(prefix or "") + len(uri) > MAX_XML_NAME_LENGTH:
            raise XlsxAdmissionError("statement_resource_limit")

    def start(self, tag, attributes):
        self.depth += 1
        self.counter[0] += 1
        if (self.depth > MAX_XML_DEPTH or self.counter[0] > MAX_XML_ELEMENTS
                or any(len(name) > MAX_XML_NAME_LENGTH for name in (tag, *attributes))):
            raise XlsxAdmissionError("statement_resource_limit")

    def end(self, tag):
        self.depth -= 1

    def data(self, data):
        pass

    def close(self):
        pass


def _preflight_xml(source, counter):
    parser = SafeElementTree.XMLParser(target=_XmlBudgetTarget(counter),
                                      forbid_dtd=True, forbid_entities=True, forbid_external=True)
    while True:
        chunk = source.read(16 * 1024)
        if not chunk:
            break
        parser.feed(chunk)
    parser.close()


def admit_xlsx_package(content: bytes) -> None:
    """Reject malformed or allocation-hostile OOXML before openpyxl sees it.

    This validates packaging and resource bounds only. Santander recognition
    remains inside the existing importer and parser contract.
    """

    try:
        with ZipFile(BytesIO(content), "r") as archive:
            infos = archive.infolist()
            if not infos:
                raise XlsxAdmissionError("xlsx_invalid")
            if len(infos) > MAX_ARCHIVE_ENTRIES:
                raise XlsxAdmissionError("statement_resource_limit")

            seen: set[str] = set()
            for info in infos:
                if info.orig_filename != info.filename:
                    raise XlsxAdmissionError("xlsx_invalid")
                name = _validate_member_name(info.filename)
                if name in seen:
                    raise XlsxAdmissionError("xlsx_invalid")
                seen.add(name)
                if info.flag_bits & 0x1:
                    raise XlsxAdmissionError("xlsx_invalid")
                if info.compress_type not in {ZIP_STORED, ZIP_DEFLATED}:
                    raise XlsxAdmissionError("xlsx_invalid")
                if info.file_size > MAX_ENTRY_EXPANDED_BYTES:
                    raise XlsxAdmissionError("statement_resource_limit")
                if _is_macro_member(name):
                    raise XlsxAdmissionError("xlsx_invalid")

            if sum(info.file_size for info in infos) > MAX_TOTAL_EXPANDED_BYTES:
                raise XlsxAdmissionError("statement_resource_limit")
            # First pass reaches EOF for every entry, validating actual expanded
            # size and CRC without retaining extracted package data.
            total_counter = [0]
            for info in infos:
                with archive.open(info, "r") as source:
                    reader = _BoundedEntryReader(source, total_counter=total_counter)
                    while reader.read(64 * 1024):
                        pass

            if "[Content_Types].xml" not in seen:
                raise XlsxAdmissionError("xlsx_invalid")
            # A namespace URI is repeated in every expanded QName. Raw XML byte
            # limits alone do not bound that amplification inside a tree parser.
            preflight_counter = [0]
            for info in infos:
                if info.filename.endswith((".xml", ".rels")):
                    with archive.open(info) as source:
                        _preflight_xml(source, preflight_counter)
            element_counter = [0]
            declared_worksheets = _declared_worksheet_names(archive, element_counter)
            if not declared_worksheets or any(name not in seen for name in declared_worksheets):
                raise XlsxAdmissionError("xlsx_invalid")
            if len(declared_worksheets) > MAX_WORKSHEETS:
                raise XlsxAdmissionError("statement_resource_limit")

            total_rectangular_cells = 0
            actual_worksheets = set()
            references = _PackageReferences()
            for info in infos:
                name = info.filename
                if not (name.endswith(".xml") or name.endswith(".rels")):
                    continue
                if name == "[Content_Types].xml":
                    continue
                # Inspect every XML part as a possible worksheet so a hostile
                # relationship cannot route openpyxl around path/content-type
                # based limits.
                extent = _SheetExtent()
                with archive.open(info, "r") as source:
                    is_worksheet = _inspect_xml(
                        source,
                        extent=extent,
                        element_counter=element_counter,
                        references=references,
                        member_name=name,
                    )
                if is_worksheet:
                    actual_worksheets.add(name)
                    total_rectangular_cells += extent.max_row * extent.max_column
                    if total_rectangular_cells > MAX_RECTANGULAR_CELLS:
                        raise XlsxAdmissionError("statement_resource_limit")

            if len(actual_worksheets) > MAX_WORKSHEETS:
                raise XlsxAdmissionError("statement_resource_limit")
            references.validate()
            if not set(references.worksheet_targets).issubset(actual_worksheets):
                raise XlsxAdmissionError("xlsx_invalid")
            if "xl/workbook.xml" not in seen or not declared_worksheets.issubset(actual_worksheets):
                raise XlsxAdmissionError("xlsx_invalid")
    except XlsxAdmissionError:
        raise
    except (BadZipFile, DefusedXmlException, OSError, RuntimeError, ValueError, EOFError):
        raise XlsxAdmissionError("xlsx_invalid") from None
    except Exception:
        raise XlsxAdmissionError("xlsx_invalid") from None


def _validate_member_name(value: str) -> str:
    if not isinstance(value, str) or not value or "\x00" in value or "\\" in value:
        raise XlsxAdmissionError("xlsx_invalid")
    path = PurePosixPath(value)
    if path.is_absolute() or any(part in {"", ".", ".."} for part in path.parts):
        raise XlsxAdmissionError("xlsx_invalid")
    normalized = path.as_posix()
    if normalized != value.removesuffix("/"):
        raise XlsxAdmissionError("xlsx_invalid")
    return normalized


def _is_macro_member(name: str) -> bool:
    lower = name.lower()
    return (
        lower.endswith(".bin")
        or "vbaproject" in lower
        or lower.startswith("xl/embeddings/")
        or lower.startswith("xl/activeX/".lower())
    )


def _declared_worksheet_names(archive: ZipFile, element_counter: list[int]) -> set[str]:
    result = set()
    depth = 0
    root_name = None
    with archive.open("[Content_Types].xml", "r") as source:
        parser = SafeElementTree.iterparse(
            source,
            events=("start", "end"),
            forbid_dtd=True,
            forbid_entities=True,
            forbid_external=True,
        )
        for event, element in parser:
            local = element.tag.rsplit("}", 1)[-1] if isinstance(element.tag, str) else ""
            if event == "end":
                depth -= 1
                element.clear()
                continue
            depth += 1
            if depth == 1:
                root_name = local
            element_counter[0] += 1
            if depth > MAX_XML_DEPTH or element_counter[0] > MAX_XML_ELEMENTS:
                raise XlsxAdmissionError("statement_resource_limit")
            if local == "Relationship" and element.attrib.get("TargetMode", "").lower() == "external":
                raise XlsxAdmissionError("xlsx_invalid")
            if local in {"Default", "Override"}:
                content_type = element.attrib.get("ContentType", "").lower()
                if "macroenabled" in content_type or "vba" in content_type:
                    raise XlsxAdmissionError("xlsx_invalid")
            if local != "Override":
                continue
            if not element.attrib.get("PartName", "").endswith(".xml"):
                raise XlsxAdmissionError("xlsx_invalid")
            if content_type == "application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml":
                part_name = element.attrib.get("PartName", "")
                if not part_name.startswith("/"):
                    raise XlsxAdmissionError("xlsx_invalid")
                result.add(_validate_member_name(part_name[1:]))
    if depth != 0 or root_name != "Types":
        raise XlsxAdmissionError("xlsx_invalid")
    return result


def _inspect_xml(source, *, extent, element_counter, references, member_name) -> bool:
    depth = 0
    current_row = 0
    current_column = 0
    root_name = None
    range_work = 0
    shared_string_size = 0
    cell_type = None
    cell_xfs_depth = None
    parser = SafeElementTree.iterparse(
        source,
        events=("start", "end"),
        forbid_dtd=True,
        forbid_entities=True,
        forbid_external=True,
    )
    for event, element in parser:
        local = element.tag.rsplit("}", 1)[-1] if isinstance(element.tag, str) else ""
        if event == "start":
            depth += 1
            if depth == 1:
                root_name = local
                references.string_tables += local == "sst"
                references.style_tables += local == "styleSheet"
            element_counter[0] += 1
            if depth > MAX_XML_DEPTH or element_counter[0] > MAX_XML_ELEMENTS:
                raise XlsxAdmissionError("statement_resource_limit")
            if local == "Relationship" and element.attrib.get("TargetMode", "").lower() == "external":
                raise XlsxAdmissionError("xlsx_invalid")
            if local == "Relationship":
                kind = element.attrib.get("Type")
                target = element.attrib.get("Target", "")
                if kind not in _SAFE_RELATIONSHIPS or not target.endswith(".xml"):
                    raise XlsxAdmissionError("xlsx_invalid")
                owner = posixpath.dirname(member_name).removesuffix("_rels")
                resolved = _validate_member_name(posixpath.normpath(posixpath.join(owner, target)).lstrip("/"))
                identifier = element.attrib.get("Id", "")
                rels = references.relationships.setdefault(member_name, {})
                if not identifier or identifier in rels:
                    raise XlsxAdmissionError("xlsx_invalid")
                rels[identifier] = (kind, resolved)
                if kind.endswith("/worksheet"):
                    references.worksheet_targets.append(resolved)
            if root_name == "workbook" and local == "sheet":
                identifier = element.attrib.get(
                    "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id", "")
                references.sheets.append(identifier)
                references.workbook_sheets.setdefault(member_name, []).append(identifier)
                if len(references.sheets) > MAX_WORKSHEETS:
                    raise XlsxAdmissionError("statement_resource_limit")
            if root_name == "sst" and local == "si":
                shared_string_size = 0
            if root_name == "styleSheet":
                if local == "numFmt":
                    index = _nonnegative_index(element.attrib.get("numFmtId", ""))
                    references.formats[index] = max(references.formats.get(index, 0), len(element.attrib.get("formatCode", "")))
                elif local == "cellXfs":
                    cell_xfs_depth = depth
                elif local == "xf" and cell_xfs_depth is not None and depth == cell_xfs_depth + 1:
                    references.styles.append(_nonnegative_index(element.attrib.get("numFmtId", "0")))
            if extent is not None and root_name == "worksheet":
                if local == "row":
                    supplied = element.attrib.get("r")
                    current_row = _positive_index(supplied) if supplied is not None else current_row + 1
                    current_column = 0
                    extent.include(current_row, max(current_column, 1))
                elif local == "c":
                    cell_type = element.attrib.get("t")
                    style = _nonnegative_index(element.attrib.get("s", "0"))
                    references.style_uses[style] = references.style_uses.get(style, 0) + 1
                    reference = element.attrib.get("r")
                    if reference is None:
                        current_column += 1
                        extent.include(max(current_row, 1), current_column)
                    else:
                        row, column = _cell_extent(reference)
                        current_row = max(current_row, row)
                        current_column = column
                        extent.include(row, column)
                elif local in {"dimension", "mergeCell", "hyperlink"}:
                    reference = element.attrib.get("ref")
                    if not reference:
                        raise XlsxAdmissionError("xlsx_invalid")
                    row, column = _range_extent(reference)
                    extent.include(row, column)
                    if local != "dimension":
                        # Repeated overlapping ranges also multiply loader work.
                        range_work += row * column
                        if range_work > MAX_RECTANGULAR_CELLS:
                            raise XlsxAdmissionError("statement_resource_limit")
        else:
            if root_name == "sst":
                if local == "t":
                    shared_string_size += len(element.text or "")
                elif local == "si":
                    references.shared_strings.append(shared_string_size)
            elif root_name == "worksheet":
                if local == "v" and cell_type == "s":
                    index = _nonnegative_index(element.text or "")
                    references.shared_uses[index] = references.shared_uses.get(index, 0) + 1
                elif local in {"t", "v", "f"}:
                    references.inline_text += len(element.text or "")
                elif local == "c":
                    cell_type = None
            if local == "cellXfs":
                cell_xfs_depth = None
            depth -= 1
            element.clear()
    if depth != 0:
        raise XlsxAdmissionError("xlsx_invalid")
    return root_name == "worksheet"


def _nonnegative_index(value):
    if not value.isascii() or not value.isdecimal() or len(value) > 10:
        raise XlsxAdmissionError("xlsx_invalid")
    return int(value)


def _positive_index(value: str) -> int:
    if not value.isascii() or not value.isdecimal() or value.startswith("0"):
        raise XlsxAdmissionError("xlsx_invalid")
    return int(value)


def _column_number(value: str) -> int:
    result = 0
    for character in value:
        result = result * 26 + ord(character) - ord("A") + 1
    return result


def _cell_extent(value: str) -> tuple[int, int]:
    match = _CELL_REFERENCE.fullmatch(value)
    if match is None:
        raise XlsxAdmissionError("xlsx_invalid")
    return int(match.group(2)), _column_number(match.group(1))


def _range_extent(value: str) -> tuple[int, int]:
    match = _RANGE_REFERENCE.fullmatch(value)
    if match is None:
        raise XlsxAdmissionError("xlsx_invalid")
    row = max(int(match.group(2)), int(match.group(4) or match.group(2)))
    column = max(_column_number(match.group(1)), _column_number(match.group(3) or match.group(1)))
    return row, column
